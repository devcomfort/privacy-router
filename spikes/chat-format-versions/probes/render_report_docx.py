from __future__ import annotations

import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).parents[3]
MAIN_REPORT = ROOT / "docs" / "message-format-report.md"
HASH_SPEC = ROOT / "docs" / "cache-hash-utility.md"
OUTPUT = ROOT / "docs" / "message-format-report.docx"


LINK_RE = re.compile(r"\[([^]]+)\]\(([^)]+)\)")
INLINE_RE = re.compile(r"(\[[^]]+\]\([^)]+\)|`[^`]+`|\*\*[^*]+\*\*)")
TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            element = borders.find(qn(f"w:{edge}"))
            if element is None:
                element = OxmlElement(f"w:{edge}")
                borders.append(element)
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), kwargs[edge])
            element.set(qn("w:color"), "D9E2F3")


def set_repeat_table_header(row) -> None:
    tr_properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_properties.append(repeat)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text: str) -> None:
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        link = LINK_RE.fullmatch(token)
        if link:
            add_hyperlink(paragraph, link.group(1), link.group(2))
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(9)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_rich_paragraph(document: Document, text: str, style: str = "Normal", *, indent: int = 0):
    paragraph = document.add_paragraph(style=style)
    if indent:
        paragraph.paragraph_format.left_indent = Cm(0.5 * indent)
    add_inline(paragraph, text)
    return paragraph


def parse_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        body = line.strip()
        if body.startswith("|"):
            body = body[1:]
        if body.endswith("|"):
            body = body[:-1]
        body = body.replace(r"\|", "\x00")
        rows.append([cell.replace("\x00", "|").strip() for cell in body.split("|")])
    return rows


def add_table(document: Document, lines: list[str]) -> None:
    rows = parse_table(lines)
    if len(rows) < 2:
        return
    if TABLE_SEPARATOR_RE.match("|" + "|".join(rows[1]) + "|"):
        rows.pop(1)
    column_count = len(rows[0])
    if any(len(row) != column_count for row in rows):
        raise ValueError(f"Markdown table column mismatch: {len(rows[0])} vs {sorted({len(row) for row in rows})}")
    table = document.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        if row_index == 0:
            set_repeat_table_header(row)
        for column_index in range(column_count):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell, top="4", left="4", bottom="4", right="4")
            if row_index == 0:
                set_cell_shading(cell, "1F4E78")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F4F7FB")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, values[column_index] if column_index < len(values) else "")
            for run in paragraph.runs:
                run.font.size = Pt(8.2)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_block(document: Document, lines: list[str], language: str | None = None) -> None:
    for index, line in enumerate(lines):
        paragraph = document.add_paragraph(style="Code")
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.space_after = Pt(0)
        if index == 0 and language:
            paragraph.paragraph_format.space_before = Pt(4)
        run = paragraph.add_run(line if line else " ")
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(8.2)
        properties = paragraph._p.get_or_add_pPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F2F2F2")
        properties.append(shading)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def render_markdown(document: Document, path: Path, *, skip_first_h1: bool = True) -> None:
    lines = path.read_text().splitlines()
    index = 0
    first_h1_skipped = not skip_first_h1
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            language = line[3:].strip() or None
            index += 1
            code: list[str] = []
            while index < len(lines) and not lines[index].startswith("```"):
                code.append(lines[index])
                index += 1
            if index < len(lines):
                index += 1
            add_code_block(document, code, language)
            continue
        if line.startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, table_lines)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1 and not first_h1_skipped:
                first_h1_skipped = True
                index += 1
                continue
            if level == 1 and skip_first_h1 and not first_h1_skipped:
                first_h1_skipped = True
                index += 1
                continue
            style_level = min(max(level - 1, 1), 4)
            paragraph = document.add_paragraph(style=f"Heading {style_level}")
            add_inline(paragraph, text)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        bullet = re.match(r"^(\s*)[-*]\s+(.+)$", line)
        if bullet:
            indent = min(len(bullet.group(1)) // 2, 2)
            paragraph = document.add_paragraph(style="List Bullet" if indent == 0 else f"List Bullet {indent + 1}")
            add_inline(paragraph, bullet.group(2))
            index += 1
            continue
        numbered = re.match(r"^(\s*)\d+\.\s+(.+)$", line)
        if numbered:
            indent = min(len(numbered.group(1)) // 2, 2)
            paragraph = document.add_paragraph(style="List Number" if indent == 0 else f"List Number {indent + 1}")
            add_inline(paragraph, numbered.group(2))
            index += 1
            continue
        if line.startswith("> "):
            paragraph = add_rich_paragraph(document, line[2:], "Quote")
            paragraph.paragraph_format.left_indent = Cm(0.6)
            index += 1
            continue
        if line.strip() == "---":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "B7C9DD")
            border.append(bottom)
            paragraph._p.get_or_add_pPr().append(border)
            index += 1
            continue
        add_rich_paragraph(document, line)
        index += 1


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for level, size, color in ((1, 16, "1F4E78"), (2, 13, "2F5597"), (3, 11, "1F4E78"), (4, 10, "5B6573")):
        style = styles[f"Heading {level}"]
        style.font.name = "맑은 고딕"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12 if level <= 2 else 8)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    if "Code" not in [style.name for style in styles]:
        code = styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code"]
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code.font.size = Pt(8.2)
    code.paragraph_format.left_indent = Cm(0.25)
    code.paragraph_format.right_indent = Cm(0.25)
    code.paragraph_format.line_spacing = 1.0

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("Privacy Router  |  채팅 메시지 형식 조사")
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("인쇄용 통합본  ·  ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)
    add_page_number(footer)


def markdown_headings(path: Path) -> list[tuple[int, str]]:
    headings: list[tuple[int, str]] = []
    in_fence = False
    for line in path.read_text().splitlines():
        if line.strip().startswith("```"):
            in_fence = not in_fence
            continue
        if in_fence:
            continue
        match = re.match(r"^(#{1,3})\s+(.+?)\s*$", line)
        if match:
            level = len(match.group(1))
            title = re.sub(r"`([^`]+)`", r"\1", match.group(2))
            title = re.sub(r"\*\*([^*]+)\*\*", r"\1", title)
            headings.append((level, title))
    return headings

def add_toc(document: Document) -> None:
    heading = document.add_paragraph(style="Heading 1")
    heading.add_run("목차")
    note = document.add_paragraph("인쇄용 정적 목차입니다. 본문 제목과 부록 구조를 기준으로 작성했습니다.")
    note.runs[0].italic = True
    note.runs[0].font.color.rgb = RGBColor(100, 100, 100)

    sources = [(MAIN_REPORT, True), (HASH_SPEC, False)]
    for source, skip_first_h1 in sources:
        entries = markdown_headings(source)
        if skip_first_h1 and entries and entries[0][0] == 1:
            entries = entries[1:]
        for level, title in entries:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.45 * max(level - 1, 0))
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run(title)
            run.font.size = Pt(9.5 if level == 3 else 10)
            if level == 1:
                run.bold = True


def add_title_page(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Cm(3.0)
    run = paragraph.add_run("채팅 메시지 형식·버전·provider 전송\n및 hash 설계 조사 보고서")
    run.bold = True
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(14)
    run = subtitle.add_run("OpenAI · Anthropic Claude · LiteLLM · LangChain Core")
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(89, 89, 89)

    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_before = Cm(4.0)
    run = metadata.add_run(f"인쇄용 파생본  ·  {date.today().isoformat()}\n문서 원본: docs/message-format-report.md\n부록 원본: docs/cache-hash-utility.md")
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(89, 89, 89)

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Cm(2.0)
    run = note.add_run("원본 Markdown에서 생성한 인쇄용 파생본입니다. 화면보다 인쇄·검토를 우선해 표, 코드, 참고 링크를 재배치했습니다.")
    run.font.size = Pt(9)
    run.italic = True
    run.font.color.rgb = RGBColor(100, 100, 100)


def main() -> None:
    output = Path(sys.argv[1]) if len(sys.argv) > 1 else OUTPUT
    document = Document()
    properties = document.core_properties
    properties.title = "채팅 메시지 형식·버전·provider 전송 및 hash 설계 조사 보고서"
    properties.subject = "docs/message-format-report.md와 docs/cache-hash-utility.md의 인쇄용 파생본"
    properties.author = "Privacy Router"
    properties.comments = f"Canonical sources: docs/message-format-report.md; docs/cache-hash-utility.md. Generated: {date.today().isoformat()}."
    properties.keywords = "Privacy Router, chat message schema, provider wire JSON, cache hash"
    properties.category = "Derived print document"
    configure_document(document)
    add_title_page(document)
    document.add_page_break()
    add_toc(document)
    document.add_page_break()
    main_heading = document.add_paragraph(style="Heading 1")
    main_heading.add_run("본문. 채팅 메시지 스키마·변환 코드 조사")
    render_markdown(document, MAIN_REPORT, skip_first_h1=True)

    document.add_page_break()
    appendix_heading = document.add_paragraph(style="Heading 1")
    appendix_heading.add_run("부록 A. 메시지·히스토리 hash 유틸리티 명세")
    render_markdown(document, HASH_SPEC, skip_first_h1=True)

    document.save(output)
    print(output)


if __name__ == "__main__":
    main()
